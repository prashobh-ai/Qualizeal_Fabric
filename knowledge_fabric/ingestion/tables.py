"""Spreadsheets and CSV (T41) — every sheet becomes a TYPED SQLite table plus
retrieval passages.

    tables/<doc>/<sheet>.sqlite          one table ``t``; columns typed by sampling
    data/facts.json["tables"]            {doc_id, doc_title, sheet, columns, rows, path}

* ``.csv`` / ``.tsv`` — stdlib ``csv``.
* ``.xlsx`` — ``openpyxl`` (the ``documents`` extra, MIT), imported lazily; when
  it is missing the caller gets ``ConverterUnavailableError`` naming the extra —
  never an empty document.

Header detection: the first row is the header when every cell is a non-numeric,
non-empty string and either a numeric value appears below it or its names are
unique. Column types: ``int`` when every non-empty sample parses as an integer,
``real`` when every one parses as a number, else ``text``. Empty cells are NULL.

Passages: one per row (``"<col>: <val>; …"``, CELL coordinate) capped by
``KF_TABLE_ROW_PASSAGES`` (default 2000), plus one sheet summary (rows,
columns, min/max per numeric column). The SQLite table is the source of truth
for arithmetic — an agent queries it instead of counting passages.
"""

from __future__ import annotations

import csv
import io
import os
import re
import sqlite3
import time

from .. import fabric_data as fd
from ..adapters.converter import ConverterUnavailableError
from ..contracts.types import ConvertedDocument, Coordinate, CoordinateKind, Region

_INT = re.compile(r"^[+-]?\d{1,18}$")
_NUM = re.compile(r"^[+-]?(\d[\d,]*\.?\d*|\.\d+)([eE][+-]?\d+)?%?$")
_SAMPLE = 500
_DEFAULT_ROW_PASSAGES = 2000


# ---------------------------------------------------------------------------
# cell typing
# ---------------------------------------------------------------------------
def is_number(value) -> bool:
    if isinstance(value, bool):
        return False
    if isinstance(value, (int, float)):
        return True
    s = str(value).strip()
    return bool(s) and bool(_NUM.match(s))


def _as_int(value):
    if isinstance(value, bool):
        raise ValueError
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        if value.is_integer():
            return int(value)
        raise ValueError
    s = str(value).strip().replace(",", "")
    if _INT.match(s):
        return int(s)
    raise ValueError


def _as_real(value) -> float:
    if isinstance(value, bool):
        raise ValueError
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip().replace(",", "")
    if s.endswith("%"):
        return float(s[:-1]) / 100.0
    return float(s)


def _empty(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def column_type(values) -> str:
    """``int`` | ``real`` | ``text`` by sampling up to ``_SAMPLE`` non-empty cells."""
    seen = 0
    all_int = True
    all_num = True
    for v in values:
        if _empty(v):
            continue
        seen += 1
        if all_int:
            try:
                _as_int(v)
            except (ValueError, TypeError):
                all_int = False
        if all_num:
            try:
                _as_real(v)
            except (ValueError, TypeError):
                all_num = False
        if not all_num or seen >= _SAMPLE:
            break
    if seen == 0:
        return "text"
    if all_int:
        return "int"
    if all_num:
        return "real"
    return "text"


def coerce(value, ctype: str):
    if _empty(value):
        return None
    try:
        if ctype == "int":
            return _as_int(value)
        if ctype == "real":
            return _as_real(value)
    except (ValueError, TypeError):
        return str(value).strip()
    return value if isinstance(value, str) else str(value)


# ---------------------------------------------------------------------------
# header + column names
# ---------------------------------------------------------------------------
def _clean_name(name, index: int) -> str:
    s = re.sub(r"\s+", " ", str(name if name is not None else "")).strip()
    return s or f"col{index + 1}"


def unique_names(names: list) -> list[str]:
    out: list[str] = []
    seen: dict[str, int] = {}
    for i, n in enumerate(names):
        base = _clean_name(n, i)
        key = base.lower()
        if key in seen:
            seen[key] += 1
            base = f"{base}_{seen[key]}"
        else:
            seen[key] = 1
        out.append(base)
    return out


def detect_header(rows: list[list]) -> tuple[list[str], list[list], bool]:
    """``(columns, body_rows, had_header)``. The first row is the header when
    all of its cells are non-empty non-numeric strings and (a) a numeric value
    appears in the rows below it or (b) its cells are unique names. A sheet
    whose first row already carries data gets ``col1..colN``."""
    rows = [list(r) for r in rows if any(not _empty(c) for c in r)]
    if not rows:
        return [], [], False
    width = max(len(r) for r in rows)
    rows = [r + [None] * (width - len(r)) for r in rows]
    first = rows[0]
    textual = all(not _empty(c) and not is_number(c) and isinstance(c, str) for c in first)
    if textual:
        below_numeric = any(is_number(c) for r in rows[1:] for c in r)
        unique = len({str(c).strip().lower() for c in first}) == len(first)
        if below_numeric or unique:
            return unique_names(first), rows[1:], True
    return [f"col{i + 1}" for i in range(width)], rows, False


def infer_columns(columns: list[str], body: list[list]) -> list[dict]:
    return [
        {"name": name, "type": column_type(r[i] for r in body)} for i, name in enumerate(columns)
    ]


# ---------------------------------------------------------------------------
# readers
# ---------------------------------------------------------------------------
def sheets_from_csv(text: str, delimiter: str, sheet: str = "sheet1") -> dict[str, list[list]]:
    rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    return {sheet: rows}


def sheets_from_xlsx(data: bytes) -> dict[str, list[list]]:
    """Every worksheet → rows of Python values (``openpyxl``, ``data_only`` so
    formulas yield their cached results). Raises ``ConverterUnavailableError``
    when the ``documents`` extra is not installed."""
    try:
        import openpyxl
    except ImportError as e:
        raise ConverterUnavailableError(
            "Excel (.xlsx) conversion needs openpyxl: install the documents extra "
            "(pip install 'qualizeal-knowledge-fabric[documents]')"
        ) from e
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out: dict[str, list[list]] = {}
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append([_cell_value(c) for c in row])
        out[ws.title] = rows
    return out


def _cell_value(v):
    if v is None or isinstance(v, (str, int, float, bool)):
        return v
    if hasattr(v, "isoformat"):
        return v.isoformat()
    return str(v)


# ---------------------------------------------------------------------------
# SQLite + facts
# ---------------------------------------------------------------------------
_SQL_TYPE = {"int": "INTEGER", "real": "REAL", "text": "TEXT"}


def _q(name: str) -> str:
    return '"' + name.replace('"', '""') + '"'


def safe_name(name: str) -> str:
    s = re.sub(r"[^A-Za-z0-9_.-]+", "_", str(name or "")).strip("._")
    return s or "sheet"


def write_sqlite(path: str, columns: list[dict], body: list[list]) -> int:
    """(Re)write ``path`` with a single table ``t``; returns the row count."""
    os.makedirs(os.path.dirname(path), exist_ok=True)
    if os.path.exists(path):
        os.remove(path)
    conn = sqlite3.connect(path)
    try:
        ddl = ", ".join(f"{_q(c['name'])} {_SQL_TYPE[c['type']]}" for c in columns)
        conn.execute(f"CREATE TABLE t ({ddl})")
        placeholders = ", ".join("?" for _ in columns)
        conn.executemany(
            f"INSERT INTO t VALUES ({placeholders})",
            ([coerce(r[i], c["type"]) for i, c in enumerate(columns)] for r in body),
        )
        conn.commit()
        n = conn.execute("SELECT COUNT(*) FROM t").fetchone()[0]
    finally:
        conn.close()
    return int(n)


def numeric_ranges(columns: list[dict], body: list[list]) -> dict[str, dict]:
    out: dict[str, dict] = {}
    for i, c in enumerate(columns):
        if c["type"] not in ("int", "real"):
            continue
        vals = [coerce(r[i], c["type"]) for r in body]
        nums = [v for v in vals if isinstance(v, (int, float)) and not isinstance(v, bool)]
        if nums:
            out[c["name"]] = {"min": min(nums), "max": max(nums)}
    return out


def register_table(entry: dict) -> str:
    """Merge one ``{doc_id, doc_title, sheet, columns, rows, path, …}`` entry into
    ``data/facts.json["tables"]`` (same doc+sheet replaces; other keys untouched)."""
    p = fd.data_path("facts.json", mkdir=True)
    facts = fd.read_json(p, {}) or {}
    tables = [
        t
        for t in facts.get("tables", [])
        if not (t.get("doc_id") == entry["doc_id"] and t.get("sheet") == entry["sheet"])
    ]
    tables.append(entry)
    facts["tables"] = sorted(tables, key=lambda t: (t.get("doc_title", ""), t.get("sheet", "")))
    fd.write_json(p, facts)
    return p


# ---------------------------------------------------------------------------
# passages
# ---------------------------------------------------------------------------
def _fmt(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def row_text(columns: list[dict], row: list) -> str:
    parts = []
    for i, c in enumerate(columns):
        v = _fmt(row[i]) if i < len(row) else ""
        if v:
            parts.append(f"{c['name']}: {v}")
    return "; ".join(parts)


def summary_text(
    sheet: str, title: str, columns: list[dict], rows: int, ranges: dict, capped: int | None
) -> str:
    names = ", ".join(c["name"] for c in columns)
    text = f"Sheet '{sheet}' of {title}: {rows} rows × {len(columns)} columns ({names})."
    if ranges:
        spans = "; ".join(
            f"{name} min {_fmt(r['min'])} max {_fmt(r['max'])}" for name, r in ranges.items()
        )
        text += f" Numeric ranges: {spans}."
    if capped is not None:
        text += (
            f" Row passages cover the first {capped} rows; the full table is queryable in SQLite."
        )
    return text


def row_passage_cap() -> int:
    try:
        return max(0, int(os.environ.get("KF_TABLE_ROW_PASSAGES", _DEFAULT_ROW_PASSAGES)))
    except ValueError:
        return _DEFAULT_ROW_PASSAGES


def convert_sheets(
    sheets: dict[str, list[list]],
    *,
    doc_id: str,
    doc_title: str,
    language: str = "en",
    citation_url: str = "",
    persist: bool = True,
) -> ConvertedDocument:
    """Sheets → typed SQLite tables (+ facts registration) and passages."""
    regions: list[Region] = []
    cap = row_passage_cap()
    now = time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime())
    for raw_name, rows in sheets.items():
        sheet = safe_name(raw_name)
        columns_, body, had_header = detect_header(rows)
        if not columns_:
            continue
        columns = infer_columns(columns_, body)
        rel = os.path.join("tables", doc_id, f"{sheet}.sqlite")
        n_rows = len(body)
        if persist:
            path = fd.path("tables", doc_id, f"{sheet}.sqlite", mkdir=True)
            n_rows = write_sqlite(path, columns, body)
            register_table(
                {
                    "doc_id": doc_id,
                    "doc_title": doc_title,
                    "sheet": sheet,
                    "sheet_title": str(raw_name),
                    "columns": columns,
                    "rows": n_rows,
                    "path": rel,
                    "header_detected": had_header,
                    "citation_url": citation_url,
                    "as_of": now,
                }
            )
        ranges = numeric_ranges(columns, body)
        capped = cap if n_rows > cap else None
        regions.append(
            Region(
                text=summary_text(sheet, doc_title, columns, n_rows, ranges, capped),
                coordinate=Coordinate(
                    CoordinateKind.CELL,
                    {"sheet": sheet, "row": 0, "col": 0, "part": "summary", "table": rel},
                ),
            )
        )
        for r, row in enumerate(body[:cap], start=1):
            text = row_text(columns, row)
            if not text:
                continue
            regions.append(
                Region(
                    text=text,
                    coordinate=Coordinate(
                        CoordinateKind.CELL,
                        {
                            "sheet": sheet,
                            "row": r,
                            "col": 0,
                            "cols": len(columns),
                            "part": "row",
                            "table": rel,
                        },
                    ),
                )
            )
    return ConvertedDocument(language=language, regions=regions)
